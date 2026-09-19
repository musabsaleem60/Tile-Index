"""Replace sanitary catalogue and inventory from the approved workbook."""

import argparse
import json
import os
import re
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import MetaData, create_engine, delete, func, select


SHEETS = {
    "Tile Index": ("Tile Index - Korangi", "Column1"),
    "Tile Cera": ("Tile Cera - Korangi", "Stock2"),
    "DHA": ("DHA", "Stock"),
}


def clean(value):
    return "" if value is None else str(value).strip()


def normalize(value):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", clean(value).lower())).strip()


def decimal_value(value, field, sheet, row_number):
    if value is None or clean(value) == "":
        if field == "stock":
            return Decimal(0)
        raise ValueError(f"{sheet} row {row_number}: blank {field}")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = clean(value).replace("Rs.", "").replace("Rs", "").replace(",", "").strip()
    try:
        return Decimal(text)
    except Exception as exc:
        raise ValueError(f"{sheet} row {row_number}: invalid {field} {value!r}") from exc


def parse_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    products = {}
    inventory = defaultdict(list)
    raw_rows = 0
    for sheet_name, (_branch_name, stock_column) in SHEETS.items():
        sheet = workbook[sheet_name]
        headers = [clean(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        columns = {header: index for index, header in enumerate(headers)}
        required = {"Company", "Product Name/Code", "Color", "Price", stock_column}
        missing_headers = required - columns.keys()
        if missing_headers:
            raise ValueError(f"{sheet_name}: missing columns {sorted(missing_headers)}")
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            company = clean(values[columns["Company"]])
            product_name = clean(values[columns["Product Name/Code"]])
            color = clean(values[columns["Color"]])
            if not any((company, product_name, color)):
                continue
            if not all((company, product_name, color)):
                raise ValueError(f"{sheet_name} row {row_number}: missing required identity field")
            raw_rows += 1
            key = (normalize(company), normalize(product_name), normalize(color))
            price = decimal_value(values[columns["Price"]], "price", sheet_name, row_number)
            stock = decimal_value(values[columns[stock_column]], "stock", sheet_name, row_number)
            if stock != stock.to_integral_value() or stock < 0:
                raise ValueError(f"{sheet_name} row {row_number}: stock must be a nonnegative integer")
            if key not in products:
                products[key] = {
                    "company_name": company,
                    "product_category": product_name,
                    "color": color,
                    "sale_price": price,
                }
            else:
                products[key]["sale_price"] = max(products[key]["sale_price"], price)
            inventory[(key, sheet_name)].append(int(stock))

    resolved_inventory = {
        key: max(values) for key, values in inventory.items()
    }
    if raw_rows != 370 or len(products) != 211 or len(resolved_inventory) != 362:
        raise ValueError(
            f"Workbook shape changed: rows={raw_rows}, products={len(products)}, inventory={len(resolved_inventory)}"
        )
    return products, resolved_inventory


def count(connection, table, where=None):
    statement = select(func.count()).select_from(table)
    if where is not None:
        statement = statement.where(where)
    return int(connection.scalar(statement))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--commit", action="store_true")
    args = parser.parse_args()
    database_url = os.environ.get("DATABASE_URL")
    if not database_url:
        raise SystemExit("DATABASE_URL is not set")

    products, inventory = parse_workbook(args.workbook)
    engine = create_engine(database_url)
    metadata = MetaData()
    metadata.reflect(bind=engine)
    sanitary_products = metadata.tables["sanitary_products"]
    sanitary_inventory = metadata.tables["sanitary_inventory"]
    sanitary_transactions = metadata.tables["sanitary_stock_transactions"]
    invoice_items = metadata.tables["invoice_items"]
    stock_transactions = metadata.tables["stock_transactions"]
    branches = metadata.tables["branches"]

    with engine.begin() as connection:
        before = {
            "sanitary_products": count(connection, sanitary_products),
            "sanitary_inventory": count(connection, sanitary_inventory),
            "sanitary_stock_transactions": count(connection, sanitary_transactions),
            "invoice_item_references": count(connection, invoice_items, invoice_items.c.sanitary_product_id.is_not(None)),
            "stock_transaction_references": count(connection, stock_transactions, stock_transactions.c.sanitary_product_id.is_not(None)),
        }
        expected = {
            "sanitary_products": 112,
            "sanitary_inventory": 1,
            "sanitary_stock_transactions": 1,
            "invoice_item_references": 0,
            "stock_transaction_references": 0,
        }
        if before != expected:
            raise RuntimeError(f"Pre-import state differs; expected {expected}, found {before}")

        branch_ids = dict(connection.execute(select(branches.c.name, branches.c.id)).all())
        missing_branches = {name for name, _ in SHEETS.values()} - branch_ids.keys()
        if missing_branches:
            raise RuntimeError(f"Missing branches: {sorted(missing_branches)}")

        if not args.commit:
            print(json.dumps({
                "mode": "dry-run",
                "before": before,
                "products_to_create": len(products),
                "inventory_to_create": len(inventory),
                "zero_inventory": sum(quantity == 0 for quantity in inventory.values()),
                "nonzero_inventory": sum(quantity > 0 for quantity in inventory.values()),
            }, indent=2))
            connection.rollback()
            return

        connection.execute(delete(sanitary_transactions))
        connection.execute(delete(sanitary_inventory))
        connection.execute(delete(sanitary_products))
        product_ids = {}
        for sequence, (key, product) in enumerate(products.items(), 1):
            result = connection.execute(sanitary_products.insert().values(
                company_name=product["company_name"],
                product_category=product["product_category"],
                color=product["color"],
                purchase_price=0,
                sale_price=float(product["sale_price"]),
                sku=f"SAN-{sequence:06d}",
            ).returning(sanitary_products.c.id))
            product_ids[key] = result.scalar_one()

        inventory_rows = []
        for (key, sheet_name), quantity in inventory.items():
            branch_name = SHEETS[sheet_name][0]
            inventory_rows.append({
                "branch_id": branch_ids[branch_name],
                "sanitary_product_id": product_ids[key],
                "quantity": quantity,
            })
        connection.execute(sanitary_inventory.insert(), inventory_rows)

        after = {
            "sanitary_products": count(connection, sanitary_products),
            "sanitary_inventory": count(connection, sanitary_inventory),
            "sanitary_stock_transactions": count(connection, sanitary_transactions),
        }
        if after != {
            "sanitary_products": 211,
            "sanitary_inventory": 362,
            "sanitary_stock_transactions": 0,
        }:
            raise RuntimeError(f"Post-import count mismatch: {after}")
        print(json.dumps({
            "mode": "commit",
            "before": before,
            "after": after,
            "zero_inventory": sum(quantity == 0 for quantity in inventory.values()),
            "nonzero_inventory": sum(quantity > 0 for quantity in inventory.values()),
        }, indent=2))
    engine.dispose()


if __name__ == "__main__":
    main()
